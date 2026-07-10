"""Accès PostgreSQL / CSV — lecture produits + mise à jour image_url."""

from __future__ import annotations

import csv
import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator
from urllib.parse import urlparse

import psycopg2
import psycopg2.extras

from config import CSV_FALLBACK, DATABASE_URL, ONLY_WITHOUT_PHOTO

logger = logging.getLogger(__name__)


@dataclass
class ProductRow:
    id: int
    name: str
    description: str | None
    category: str | None
    image_url: str | None
    sku: str | None = None


def _parse_dsn(url: str) -> dict:
    """Convertit DATABASE_URL postgres://… en kwargs psycopg2."""
    u = urlparse(url)
    return {
        "host": u.hostname or "localhost",
        "port": u.port or 5432,
        "dbname": (u.path or "/qdia_export").lstrip("/") or "qdia_export",
        "user": u.username or "postgres",
        "password": u.password or "",
        "sslmode": "prefer",
    }


class Database:
    def __init__(self, dsn: str | None = None) -> None:
        self.dsn = dsn or DATABASE_URL
        self._conn = None

    def connect(self) -> None:
        if not self.dsn:
            raise RuntimeError("DATABASE_URL manquant — définissez-le dans .env")
        self._conn = psycopg2.connect(**_parse_dsn(self.dsn))
        self._conn.autocommit = False
        logger.info("Connexion PostgreSQL OK")

    def close(self) -> None:
        if self._conn:
            self._conn.close()
            self._conn = None

    def __enter__(self) -> "Database":
        self.connect()
        return self

    def __exit__(self, *_) -> None:
        self.close()

    def fetch_products(
        self,
        *,
        limit: int = 0,
        only_without_photo: bool | None = None,
        skip_ids: set[int] | None = None,
    ) -> list[ProductRow]:
        only = ONLY_WITHOUT_PHOTO if only_without_photo is None else only_without_photo
        skip_ids = skip_ids or set()

        sql = """
            SELECT id, name, description, category, image_url, sku
            FROM products
            WHERE 1=1
        """
        params: list = []
        if only:
            sql += """
              AND (
                image_url IS NULL
                OR image_url = ''
                OR image_url LIKE '%%qdia-photo-placeholder%%'
                OR image_url LIKE '%%Photo IA%%'
              )
            """
        sql += " ORDER BY id ASC"
        if limit and limit > 0:
            sql += " LIMIT %s"
            params.append(limit)

        assert self._conn is not None
        with self._conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(sql, params)
            rows = cur.fetchall()

        products = [
            ProductRow(
                id=int(r["id"]),
                name=str(r["name"] or "").strip(),
                description=r.get("description"),
                category=r.get("category"),
                image_url=r.get("image_url"),
                sku=r.get("sku"),
            )
            for r in rows
            if r["name"] and int(r["id"]) not in skip_ids
        ]
        logger.info("%d produit(s) à traiter (skip=%d)", len(products), len(skip_ids))
        return products

    def update_image(
        self,
        product_id: int,
        image_url: str,
        *,
        description: str | None = None,
    ) -> None:
        assert self._conn is not None
        with self._conn.cursor() as cur:
            if description is not None:
                cur.execute(
                    """
                    UPDATE products
                    SET image_url = %s,
                        images = ARRAY[%s]::text[],
                        description = COALESCE(NULLIF(%s, ''), description)
                    WHERE id = %s
                    """,
                    (image_url, image_url, description, product_id),
                )
            else:
                cur.execute(
                    """
                    UPDATE products
                    SET image_url = %s,
                        images = ARRAY[%s]::text[]
                    WHERE id = %s
                    """,
                    (image_url, image_url, product_id),
                )
        self._conn.commit()


def load_from_csv(path: str | Path) -> list[ProductRow]:
    """Fallback CSV/Excel (colonnes: id, name, description, category, image_url)."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)

    if path.suffix.lower() in (".xlsx", ".xls"):
        import openpyxl  # type: ignore

        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb.active
        headers = [str(c.value or "").strip().lower() for c in next(ws.iter_rows(max_row=1))]
        products: list[ProductRow] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row))
            if not data.get("name"):
                continue
            products.append(
                ProductRow(
                    id=int(data.get("id") or 0),
                    name=str(data["name"]),
                    description=str(data.get("description") or "") or None,
                    category=str(data.get("category") or "") or None,
                    image_url=str(data.get("image_url") or "") or None,
                    sku=str(data.get("sku") or "") or None,
                )
            )
        return products

    products = []
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            name = (row.get("name") or "").strip()
            if not name:
                continue
            products.append(
                ProductRow(
                    id=int(row.get("id") or 0),
                    name=name,
                    description=(row.get("description") or None),
                    category=(row.get("category") or None),
                    image_url=(row.get("image_url") or None) or None,
                    sku=(row.get("sku") or None),
                )
            )
    return products


def iter_products(
    *,
    limit: int = 0,
    skip_ids: set[int] | None = None,
) -> tuple[list[ProductRow], Database | None]:
    """Charge depuis Postgres, sinon CSV_FALLBACK."""
    if DATABASE_URL:
        db = Database()
        db.connect()
        return db.fetch_products(limit=limit, skip_ids=skip_ids), db
    if CSV_FALLBACK:
        rows = load_from_csv(CSV_FALLBACK)
        skip = skip_ids or set()
        rows = [r for r in rows if r.id not in skip]
        if limit:
            rows = rows[:limit]
        return rows, None
    raise RuntimeError("Ni DATABASE_URL ni CSV_FALLBACK configuré")
